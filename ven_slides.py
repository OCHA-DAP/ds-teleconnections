"""Venezuela El Niño 2026/27: a short slide deck from the same analysis as the page (ven_deep_dive.py).

    uv run --with python-pptx python ven_slides.py [--publish]

Writes out/ven/slides/: simple slide figures, content.json (every number on the slides, from the build) and an editable
PowerPoint, venezuela-el-nino-2026-27.pptx (native text, placeholders and a native chart, so it can be edited). Where LibreOffice
is installed it also renders the deck to PDF and one PNG per slide (pdftoppm), for checking and for the slides page.
--publish copies the .pptx, the PNGs and a page showing them to docs/enso/ven/slides/.

The same deck also lives as a Claude Slides artifact, for comments; its slide files are written by hand from content.json, so
re-check them against it after a rebuild.
"""
from __future__ import annotations

import argparse
import html
import json
import shutil
import subprocess
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib import colors as mcolors
from matplotlib.patches import Patch

import ven_deep_dive as V

OUT = V.OUT / "slides"
DOCS = V.DOCS / "slides"
PPTX = "venezuela-el-nino-2026-27.pptx"
PAGE_URL = "https://ocha-dap.github.io/ds-teleconnections/enso/ven/"
FEWS_VEN = "https://fews.net/latin-america-and-caribbean/venezuela"
DARK, LIGHT, PAPER, INK, MUTED, ACCENT, PANEL = "3A2414", "FBF7F2", "FBF8F4", "1F2328", "5B6168", "B5541B", "F1ECE5"
VEN_XLIM, VEN_YLIM = (-73.6, -59.6), (0.6, 12.4)
plt.rcParams.update({"font.size": 13, "axes.titlesize": 15})


def _save(fig, name: str) -> Path:
    p = OUT / f"{name}.png"
    fig.savefig(p, dpi=200, bbox_inches="tight", facecolor="#ffffff")
    plt.close(fig)
    return p


def _ven_axes(ax):
    ax.set_xlim(*VEN_XLIM)
    ax.set_ylim(*VEN_YLIM)
    ax.set_aspect("equal")
    ax.set_axis_off()


# --------------------------------------------------------------------------- figures: large type, one idea each
def fig_zone_map() -> Path:
    z, st = V.zones(), V.states()
    fig, ax = plt.subplots(figsize=(6.4, 5.6))
    z.plot(ax=ax, color=z.zone.map(V.ZONE_FILL), edgecolor="none")
    st.boundary.plot(ax=ax, color="#ffffff", lw=0.5)
    for _, row in z.iterrows():
        pt = max(getattr(row.geometry, "geoms", [row.geometry]), key=lambda q: q.area).representative_point()
        ax.text(pt.x, pt.y, row.zone, ha="center", va="center", fontsize=20, fontweight="bold", color="#0b0b0b")
    _ven_axes(ax)
    ax.legend(handles=[Patch(color=V.ZONE_FILL[k], label=f"{k}: rains start {V.ZONE_ONSET_TXT[k][0]}") for k in V.ZONE_ONSET],
              loc="lower left", fontsize=11, frameon=False)
    return _save(fig, "zone_map")


def fig_vegnow() -> Path:
    """Cropland only (the page also maps rangeland)."""
    st = V.states()
    lookup = {V._norm(n): pc for pc, n in zip(st.pcode, st.name)}
    a = V.asap(220, 1, 2)
    z = a[a.date == a.date.max()].assign(pcode=lambda d: d.region_name.map(lambda n: lookup.get(V._norm(n)))).set_index("pcode")["value"]
    zb = [-3, -1.5, -1, -0.5, 0.5, 1, 1.5, 3]
    cm, nm = mcolors.ListedColormap(V.DIV), mcolors.BoundaryNorm(zb, len(V.DIV))
    fig, ax = plt.subplots(figsize=(6.4, 5.9))
    st.merge(z.rename("z"), left_on="pcode", right_index=True, how="left").plot(
        ax=ax, column="z", cmap=cm, norm=nm, edgecolor="#ffffff", lw=0.6, missing_kwds=dict(color=V.DRY_GREY))
    _ven_axes(ax)
    cb = fig.colorbar(plt.cm.ScalarMappable(norm=nm, cmap=cm), ax=ax, orientation="horizontal", fraction=0.05, pad=0.02)
    cb.set_ticks(zb[1:-1], labels=[f"{t:g}".replace("-", "−") for t in zb[1:-1]])
    cb.set_label("cropland vegetation, z-score (warning level −1)", fontsize=12)
    return _save(fig, "vegnow")


def fig_guri() -> Path:
    g = V.guri()
    gy = V.guri_years(V.phases())
    drops = gy.loc[gy.phase == "nino", "drop"].dropna()
    peak = float(g[f"{V.NOW}-08-01":].max())
    lo, hi = peak - drops.max(), peak - drops.min()
    fig, ax = plt.subplots(figsize=(10, 4.6))
    s = g["2008-01-01":]
    ax.plot(s.index, s.values, color="#3b3a37", lw=1.3)
    low10 = float(gy.loc[2009, "low"])
    ax.axhline(low10, color=V.NINO_C, lw=1.2, ls=(0, (4, 3)))
    ax.text(pd.Timestamp("2008-03-01"), low10 - 0.5, f"2010 low, power rationing: {low10:.0f} m", color=V.NINO_C,
            fontsize=12, va="top")
    x = pd.Timestamp(V.NOW + 1, 5, 1)
    ax.plot([x, x], [lo, hi], color=V.NOW_C, lw=6, solid_capstyle="butt")
    ax.text(x - pd.Timedelta(days=60), low10 - 0.9, f"May {V.NOW + 1}, if it falls as\nafter past El Niño years:\n{lo:.0f}–{hi:.0f} m",
            ha="right", va="top", fontsize=12, fontweight="bold")   # under the 2010 line, where 2018–2026 leave room
    ax.set_ylabel("Guri reservoir level, m (satellite)")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color=V.GRID, lw=0.6)
    ax.set_xlim(pd.Timestamp("2008-01-01"), pd.Timestamp(V.NOW + 1, 12, 1))
    return _save(fig, "guri")


def fig_forecast() -> Path:
    """October–February only, the months the page's headline is about; same classes and colours as the page."""
    pr, sk, _, _ = V.seas5_rank()
    rainy = V.seas5_rainy()
    st = V.states()
    fig, axes = plt.subplots(1, 3, figsize=(12, 4.3))
    seen, mod = set(), False
    for ax, code in zip(axes, ("OND", "NDJ", "DJF")):
        cls = {pc: V.seas5_class(pr.loc[pc, code], sk.loc[pc, code], bool(rainy.loc[pc, code])) for pc in V.STATE_ORDER}
        g = st.assign(cls=st.pcode.map(lambda pc: cls[pc][0]), skill=st.pcode.map(lambda pc: cls[pc][1]))
        g.plot(ax=ax, color=g.cls.map(V.SEAS5_FILL), edgecolor="none")
        hatch = (g.skill == "mod") & (g.cls != "none")
        if hatch.any():
            g[hatch].plot(ax=ax, facecolor="none", edgecolor=V.SEAS5_HATCH_WHITE, hatch="////", lw=0)
        st.boundary.plot(ax=ax, color="#ffffff", lw=0.5)
        seen |= set(g.cls)
        mod |= bool(hatch.any())
        m = V.FC_SEASONS[code]
        ax.set_title(V._cap(V._span(m[0], m[-1])), fontsize=16, fontweight="bold")
        _ven_axes(ax)
    labels = {"vsev_dry": "Strongly below normal (driest 10%)", "sev_dry": "Below normal", "none": "Roughly normal",
              "sev_wet": "Above normal", "vsev_wet": "Strongly above normal", "off_season": "Outside the rainy season"}
    handles = [Patch(color=V.SEAS5_FILL[k], label=lab) for k, lab in labels.items() if k in seen]
    if mod:
        handles.append(Patch(facecolor=V.SEAS5_FILL["vsev_dry"], edgecolor="#ffffff", hatch="////", label="Hatched: moderate skill"))
    fig.legend(handles=handles, loc="lower center", ncols=len(handles), fontsize=12, frameon=False, bbox_to_anchor=(0.5, -0.03))
    fig.subplots_adjust(wspace=0.02, bottom=0.1)
    return _save(fig, "forecast")


def fig_rain2026(regions: list[str], values: list[int]) -> Path:
    """For the artifact deck; the .pptx has a native chart of the same numbers."""
    fig, ax = plt.subplots(figsize=(10, 4.8))
    ax.barh(range(len(values)), values, color=V.NINO_C, height=0.62)
    for i, v in enumerate(values):
        ax.text(v + 1, i, f"{v}%", va="center", fontsize=13)
    ax.axvline(100, color="#3b3a37", lw=1.2)
    ax.text(100, len(values) - 0.45, "  normal", fontsize=12, va="bottom")
    ax.set_yticks(range(len(values)), regions, fontsize=13)
    ax.set_xlim(0, 112)
    ax.set_xlabel("% of the 1991–2020 normal")
    ax.spines[["top", "right"]].set_visible(False)
    return _save(fig, "rain2026")


# --------------------------------------------------------------------------- content: every number from the page's build
def content(v: dict, rain: pd.Series) -> dict:
    span, yr = v["now_span"], v["fc_year"]
    r = rain.sort_values()
    return {
        "title": "Venezuela: El Niño 2026/27",
        "subtitle": "What past El Niño years brought, how this year has gone so far, and the forecast for the coming months",
        "byline": f"OCHA Centre for Humanitarian Data · rain to {v['data_end']}, vegetation to {v['veg_now_date']}, "
                  f"SEAS5 issued October {yr}",
        "summary_title": "Summary",
        "summary": [
            ["1", "Past El Niño years", f"A drier dry season in the south ({v['ab_dry_lo']}–{v['ab_dry_hi']}% less rain) and weaker "
                                         f"rains in the north ({v['cde_rain_lo']}–{v['cde_rain_hi']}% less); fewer flood reports; "
                                         "the Guri reservoir fell further."],
            ["2", "This year so far", f"{span} rain was {v['now_lo']}–{v['now_hi']}% of normal, among the {v['now_rank_txt']} "
                                       f"driest since 1981. Cropland is below the warning level in {v['veg_warn']}. Guri is "
                                       "near full."],
            ["3", "The forecast", f"SEAS5 (October {yr}): October–February among the driest 10% of forecasts since 1981, "
                                   "wherever it is the rainy season. For this season, plan on the forecast."],
        ],
        "zones": {"title": "In past El Niño years, the dry season was drier in the south and the rains weaker in the north",
                  "big1": f"{v['ab_dry_lo']}–{v['ab_dry_hi']}%", "lab1": "less dry-season rain in the south (zones A and B)",
                  "big2": f"{v['cde_rain_lo']}–{v['cde_rain_hi']}%", "lab2": "less rainy-season rain in the north (zones C to E)",
                  "note": f"Medians over the {v['n_nino']} El Niño years since 1981. Zones: INAMEH’s map of rainy-season onset. "
                          "Rain: CHIRPS v3."},
        "impacts": {"title": "El Niño years also had fewer flood reports, a larger fall in the Guri reservoir and poorer "
                             "dry-season vegetation",
                    "cards": [[f"{v['di_fewer_lo']}–{v['di_fewer_hi']}%", "fewer flood and landslide reports in a typical El Niño "
                                                                          "year (DesInventar, 1975–2013)"],
                              [f"{v['guri_nino']} m", f"fall of the Guri reservoir through the dry season after a typical "
                                                      f"El Niño year, against {v['guri_nina']}\u00a0m after La Niña"],
                              ["3 poorest", f"dry seasons for vegetation since 2000 all followed El Niño ({v['veg_worst']})"]],
                    "note": "Harvests show no clear link with El Niño (FAOSTAT maize and rice yields)."},
        "rain": {"title": f"{span} 2026 rain was {v['now_lo']}–{v['now_hi']}% of normal, among the {v['now_rank_txt']} driest "
                          "since 1981",
                 "regions": list(r.index), "values": [round(float(x)) for x in r.values],
                 "note": f"CHIRPS v3 to {v['data_end']}, % of the 1991–2020 normal. ERA5 gives {v['era5_lo']}–{v['era5_hi']}%. "
                         "October data come out in November."},
        "veg": {"title": f"Mid-October: vegetation near normal nationally, cropland below the warning level in {v['veg_warn']}",
                "bullets": [f"Cropland below normal in {v['veg_now_below']} of {v['veg_now_n']} states",
                            "FEWS NET (August): localized soil-moisture deficits in rainfed areas, no major crop losses"],
                "note": f"JRC ASAP, {v['veg_now_date']}: MODIS FPAR z-score over cropland; ASAP’s warning level is −1."},
        "guri": {"title": f"The Guri reservoir is near full ({v['guri_peak']} m); falls like those after past El Niño years would "
                          f"take it to {v['guri_lo']}–{v['guri_hi']} m by May {v['next_year']}",
                 "note": f"The low end is {v['guri_gap10']} m above the 2010 low. How far it falls also depends on releases for power. "
                         "NASA G-REALM satellite altimetry."},
        "forecast": {"title": "SEAS5 forecasts one of the driest October–February periods since 1981, wherever it is the "
                              "rainy season",
                     "bullets": [f"Every rainy-season state and season ({v['fc_cells']} in all) is among the driest 10% of forecasts "
                                 "since 1981",
                                 f"The forecast has skill for these months (correlation {v['fc_r_lo']}–{v['fc_r_hi']})",
                                 "ECMWF issues a new forecast every month"],
                     "note": f"ECMWF SEAS5, issued October {yr}, detrended; classes as in the team’s seas5-skill viewer."},
        "calendar": {"title": "The forecast months cover the summer planting of maize and rice and the coffee and sugar-cane "
                              "harvest",
                     "bullets": ["FEWS NET expects Stressed (IPC Phase 2) outcomes for October 2026 to January 2027",
                                 "It expects localized rainfed losses with limited impact, given irrigation and imports",
                                 "It puts peak needs in August–October 2026, driven by the earthquake and the economy"],
                     "note": "FEWS NET: seasonal calendar (Venezuela page) and El Niño briefing, August 2026."},
        "watch": {"title": "What to watch next",
                  "bullets": ["The November SEAS5 forecast, out in early November",
                              "October rain, from CHIRPS and ERA5 in November",
                              "The Guri reservoir level through the dry season",
                              "Cropland warnings in the north-west and north-east (JRC ASAP, every ten days)"],
                  "link": PAGE_URL},
        "notes": {
            "summary": "One line for each part of the deck: what past El Niño years brought, how this year has gone, and what "
                       "the forecast says. Past years are context; the forecast is what to plan on.",
            "zones": "INAMEH’s zones by rainy-season onset run from early April in the far south (A) to late June on the "
                     "northern coast (E). Each zone’s rainy season runs from its onset to November; its dry season from "
                     "December to the onset.",
            "impacts": "Floods: DesInventar records, each year against the years around it. Guri: NASA G-REALM satellite "
                       "altimetry since 1992. Vegetation: JRC ASAP FPAR, November–April, since 2000; 2023/24, a strong El Niño, "
                       "was not among the poorest.",
            "rain": "Each region’s total for the months so far this rain year, as a share of its 1991–2020 normal.",
            "veg": "The full page also maps rangeland. ASAP updates every ten days.",
            "guri": "Satellite heights above the EGM2008 geoid, not the operator’s gauge datum: compare falls, not levels.",
            "forecast": "The October forecasts in 1997 and 2015 ranked about as dry; the rain that followed was among the driest "
                        "on record in 1997 and below normal in 2015. February–April leans wet in the west and centre, where "
                        "the forecast has low skill or the months fall outside the rainy season.",
            "calendar": f"Calendar: {FEWS_VEN}",
            "watch": "Things that will change this picture over the coming weeks.",
        },
    }


# --------------------------------------------------------------------------- PowerPoint
def _fit_layouts(prs, k: float) -> None:
    """python-pptx's default template is 4:3: stretch its master and layout placeholders to the 16:9 width, and set left-aligned
    titles and smaller body text, so slides added in PowerPoint match the deck."""
    from pptx.oxml.ns import qn

    for shp in [*prs.slide_master.placeholders, *(p for lay in prs.slide_layouts for p in lay.placeholders)]:
        if shp._element.spPr.find(qn("a:xfrm")) is not None:   # a layout placeholder without one inherits the master's
            shp.left, shp.width = int(shp.left * k), int(shp.width * k)
    styles = prs.slide_master.element.find(qn("p:txStyles"))
    t = styles.find(qn("p:titleStyle")).find(qn("a:lvl1pPr"))
    t.set("algn", "l")
    t.find(qn("a:defRPr")).set("sz", "2600")
    t.find(qn("a:defRPr")).set("b", "1")
    for lvl, sz in (("a:lvl1pPr", "2000"), ("a:lvl2pPr", "1800"), ("a:lvl3pPr", "1600")):
        styles.find(qn("p:bodyStyle")).find(qn(lvl)).find(qn("a:defRPr")).set("sz", sz)


def build_pptx(c: dict, figs: dict[str, Path], path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    rgb = RGBColor.from_string
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    _fit_layouts(prs, 13.333 / 10)
    TITLE, TITLE_ONLY = prs.slide_layouts[0], prs.slide_layouts[5]
    n = c["notes"]

    def text(slide, x, y, w, h, lines, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, name=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.space_after = Pt(6)
            r = para.add_run()
            r.text = line
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, rgb(color)
        return tb

    def bullets(slide, x, y, w, h, items, size=18):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.name = "Bullets"
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pPr = para._p.get_or_add_pPr()   # a real bullet (not a typed one), with a hanging indent
            pPr.set("marL", str(Inches(0.3)))
            pPr.set("indent", str(-Inches(0.3)))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
            para.space_after = Pt(14)
            r = para.add_run()
            r.text = item
            r.font.size, r.font.color.rgb = Pt(size), rgb(INK)
        return tb

    def panel(slide, x, y, w, h, name):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = name
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(PANEL)
        box.line.fill.background()
        box.shadow.inherit = False
        box._element.find(qn("p:style")).find(qn("a:effectRef")).set("idx", "0")   # no theme shadow either
        return box

    def slide(title, eyebrow=None, note=None, notes=None):
        s = prs.slides.add_slide(TITLE_ONLY)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(PAPER)
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.75), Inches(12.1), Inches(1.15)
        t.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        t.text_frame.word_wrap = True
        t.text = title
        for r in t.text_frame.paragraphs[0].runs:
            r.font.color.rgb = rgb(INK)
        if eyebrow:
            text(s, 0.6, 0.35, 8, 0.4, eyebrow, size=14, color=ACCENT, bold=True, name="Part")
        if note:
            text(s, 0.6, 6.85, 12.1, 0.45, note, size=12, color=MUTED, name="Source")
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    PART1, PART2, PART3 = "1 · PAST EL NIÑO YEARS", "2 · THIS YEAR SO FAR", "3 · THE FORECAST"

    # 1 title
    s = prs.slides.add_slide(TITLE)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(DARK)
    for ph, top, height in ((s.shapes.title, 2.1, 1.5), (s.placeholders[1], 3.75, 1.6)):   # all four, or the box collapses
        ph.left, ph.top, ph.width, ph.height = Inches(1.0), Inches(top), Inches(11.3), Inches(height)
    s.shapes.title.text = c["title"]
    s.shapes.title.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size, r.font.color.rgb = Pt(48), rgb(LIGHT)
    sub = s.placeholders[1]
    sub.text = c["subtitle"]
    sub.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    for r in sub.text_frame.paragraphs[0].runs:
        r.font.size, r.font.color.rgb = Pt(22), rgb("E9DCCD")
    text(s, 1.0, 6.4, 11.3, 0.5, c["byline"], size=14, color="CDB9A4", name="Byline")

    # 2 summary
    s = slide(c["summary_title"], notes=n["summary"])
    for i, (num, head, body) in enumerate(c["summary"]):
        y = 2.0 + i * 1.55
        panel(s, 0.6, y, 12.1, 1.35, f"Panel {num}")
        text(s, 0.85, y + 0.2, 0.7, 0.9, num, size=40, color=ACCENT, bold=True, name=f"Number {num}")
        text(s, 1.7, y + 0.15, 3.0, 1.1, head, size=20, bold=True, name=f"Heading {num}")
        text(s, 4.6, y + 0.15, 7.9, 1.15, body, size=17, name=f"Text {num}")

    # 3 zones
    z = c["zones"]
    s = slide(z["title"], PART1, z["note"], n["zones"])
    s.shapes.add_picture(str(figs["zone_map"]), Inches(0.6), Inches(2.0), height=Inches(4.75))
    for i, (big, lab) in enumerate(((z["big1"], z["lab1"]), (z["big2"], z["lab2"]))):
        y = 2.3 + i * 2.15
        text(s, 7.0, y, 5.7, 1.0, big, size=54, color=ACCENT, bold=True, name=f"Figure {i + 1}")
        text(s, 7.0, y + 1.0, 5.7, 0.9, lab, size=20, name=f"Label {i + 1}")

    # 4 impacts
    m = c["impacts"]
    s = slide(m["title"], PART1, m["note"], n["impacts"])
    for i, (big, lab) in enumerate(m["cards"]):
        x = 0.6 + i * 4.1
        panel(s, x, 2.3, 3.85, 3.9, f"Card {i + 1}")
        text(s, x + 0.3, 2.6, 3.3, 1.0, big, size=40, color=ACCENT, bold=True, name=f"Figure {i + 1}")
        text(s, x + 0.3, 3.75, 3.3, 2.3, lab, size=18, name=f"Label {i + 1}")

    # 5 rain 2026: a native chart
    r = c["rain"]
    s = slide(r["title"], PART2, r["note"], n["rain"])
    cd = CategoryChartData()
    cd.categories = r["regions"]
    cd.add_series("% of normal", r["values"])
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(2.0), Inches(12.1), Inches(4.75), cd).chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(14)
    pl = ch.plots[0]
    pl.gap_width = 60
    pl.has_data_labels = True
    pl.data_labels.number_format, pl.data_labels.number_format_is_linked = '0"%"', False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.series[0].format.fill.solid()
    pl.series[0].format.fill.fore_color.rgb = rgb(ACCENT)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 100, 20
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("DCD8D2")
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0"%"', False
    va.format.line.fill.background()
    ch.category_axis.format.line.color.rgb = rgb("9A9993")

    # 6 vegetation now
    g = c["veg"]
    s = slide(g["title"], PART2, g["note"], n["veg"])
    s.shapes.add_picture(str(figs["vegnow"]), Inches(0.9), Inches(2.0), height=Inches(4.75))
    bullets(s, 7.0, 2.5, 5.7, 3.5, g["bullets"], size=20)

    # 7 Guri
    g = c["guri"]
    s = slide(g["title"], PART2, g["note"], n["guri"])
    s.shapes.add_picture(str(figs["guri"]), Inches(1.6), Inches(2.05), height=Inches(4.6))

    # 8 forecast
    f = c["forecast"]
    s = slide(f["title"], PART3, f["note"], n["forecast"])
    s.shapes.add_picture(str(figs["forecast"]), Inches(0.5), Inches(2.3), width=Inches(8.5))
    bullets(s, 9.3, 2.2, 3.45, 4.4, f["bullets"], size=17)

    # 9 calendar
    k = c["calendar"]
    s = slide(k["title"], PART3, k["note"], n["calendar"])
    s.shapes.add_picture(str(figs["calendar"]), Inches(0.6), Inches(2.1), width=Inches(8.2))
    bullets(s, 9.1, 2.2, 3.65, 4.4, k["bullets"], size=17)

    # 10 what to watch
    w = c["watch"]
    s = slide(w["title"], notes=n["watch"])
    bullets(s, 0.6, 2.1, 12.1, 3.8, w["bullets"], size=24)
    tb = text(s, 0.6, 6.2, 12.1, 0.5, "Full analysis, in English and Spanish: ", size=16, color=MUTED, name="Link")
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = w["link"]
    run.hyperlink.address = w["link"]
    run.font.size = Pt(16)

    prs.core_properties.title = c["title"]
    prs.core_properties.author = "OCHA Centre for Humanitarian Data"
    prs.save(path)


def render_pngs(pptx: Path) -> list[Path]:
    """LibreOffice (headless) → PDF → one PNG per slide, for checking the deck and for the slides page. Skipped, with a message,
    where LibreOffice or pdftoppm is missing. (Scripting PowerPoint or Keynote fails here: their sandbox refuses the file.)"""
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf = pptx.with_suffix(".pdf")
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)], check=True,
                       capture_output=True, timeout=300)
        for old in pptx.parent.glob("slide-*.png"):
            old.unlink()
        subprocess.run(["pdftoppm", "-png", "-r", "110", str(pdf), str(pptx.parent / "slide")], check=True, timeout=180)
    except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        print(f"no slide images ({e}); check the deck by opening it")
        return []
    return sorted(pptx.parent.glob("slide-*.png"))


PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title} — slides</title><meta name="description" content="{desc}">
<style>
:root{{--ink:#1f2328;--muted:#5b6168;--accent:#b5541b;--paper:#f4f2ee;--line:#dedad3}}
body{{margin:0;background:var(--paper);color:var(--ink);font:16px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,Arial,sans-serif}}
main{{max-width:1100px;margin:0 auto;padding:20px 16px 48px}}
a{{color:#1f5fa8}}
.top{{display:flex;flex-wrap:wrap;gap:10px;align-items:center;margin:0 0 18px}}
.btn{{display:inline-block;padding:8px 14px;border:1px solid var(--line);border-radius:6px;background:#fff;text-decoration:none;font-size:14px}}
.btn.primary{{background:var(--accent);border-color:var(--accent);color:#fff;font-weight:600}}
h1{{font-size:26px;line-height:1.25;margin:0 0 6px}}
p.meta{{color:var(--muted);margin:0 0 18px;font-size:14px}}
figure{{margin:0 0 18px;background:#fff;border:1px solid var(--line);border-radius:6px;overflow:hidden}}
figure img{{display:block;width:100%;height:auto}}
</style></head><body><main>
<div class="top"><a class="btn" href="../">← Full analysis</a><a class="btn primary" href="{pptx}" download>Download the PowerPoint (editable)</a></div>
<h1>{title}: slides</h1>
<p class="meta">{n} slides, the short version of <a href="../">the full analysis</a> (also <a href="../es/">in Spanish</a>). The
PowerPoint has editable text and a native chart; sources are on each slide and in the speaker notes.</p>
{figs}
</main></body></html>
"""


def write_page(c: dict, pngs: list[Path], dest: Path) -> None:
    figs = "\n".join(f'<figure><img src="{p.name}" alt="Slide {i}" loading="lazy"></figure>' for i, p in enumerate(pngs, 1))
    dest.write_text(PAGE.format(title=html.escape(c["title"]), desc=html.escape(c["subtitle"]), pptx=PPTX, n=len(pngs), figs=figs),
                    encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--publish", action="store_true", help="also copy the deck, its slide images and a page to docs/enso/ven/slides/")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    a = V.analyse()
    V.set_lang("en")
    _, v = V.figures(a)
    last = a["mon"].dropna(how="all").index.max()
    rain = V._region_pct(a["mon"], list(range(5, last.month + 1))).loc[V.NOW]
    c = content(v, rain)
    figs = dict(zone_map=fig_zone_map(), vegnow=fig_vegnow(), guri=fig_guri(), forecast=fig_forecast(),
                rain2026=fig_rain2026(c["rain"]["regions"], c["rain"]["values"]), calendar=V.FEWS_DIR / "ve_seasonal_calendar_rural.png")
    (OUT / "content.json").write_text(json.dumps(c, ensure_ascii=False, indent=1), encoding="utf-8")
    build_pptx(c, figs, OUT / PPTX)
    pngs = render_pngs(OUT / PPTX)
    write_page(c, pngs, OUT / "index.html")
    print(f"wrote {OUT / PPTX}, {len(pngs)} slide images and {OUT / 'index.html'}")
    if args.publish:
        if not pngs:
            raise SystemExit("no slide images: the slides page would be empty")
        if DOCS.exists():
            shutil.rmtree(DOCS)
        DOCS.mkdir(parents=True)
        for p in [OUT / PPTX, OUT / "index.html", *pngs]:
            shutil.copy(p, DOCS / p.name)
        print(f"copied the deck and its page to {DOCS}")


if __name__ == "__main__":
    main()
